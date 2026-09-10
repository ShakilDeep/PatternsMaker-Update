import zipfile
from io import BytesIO

import openpyxl

from app.domain.catalog import SIZES
from app.domain.catalog_map import map_code
from app.infrastructure.formulas import cell_value, number
from app.infrastructure.pdf_parser import parse_pdf

__all__ = ["parse_pdf", "parse_xlsx"]


def parse_xlsx(data: bytes, filename: str) -> list[dict]:
    with zipfile.ZipFile(BytesIO(data)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 30_000_000:
            raise ValueError("Workbook expanded content exceeds 30 MB")
    wb = openpyxl.load_workbook(BytesIO(data), data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(BytesIO(data), data_only=True, keep_links=False)
    result = []
    for sheet in wb:
        if sheet.max_row > 2000 or sheet.max_column > 100:
            raise ValueError("Workbook exceeds 2000 rows or 100 columns")
        header = next(
            (
                row
                for row in sheet.iter_rows(max_row=min(sheet.max_row, 30))
                if sum(str(c.value).strip() in SIZES for c in row) >= 2
            ),
            None,
        )
        if not header:
            continue
        columns = {str(c.value).strip(): c.column for c in header if str(c.value).strip() in SIZES}
        tolerance_col = next((c.column for c in header if "TOLERANCE" in str(c.value).upper()), None)
        for row in sheet.iter_rows(min_row=header[0].row + 1):
            code = str(row[0].value or "").strip()
            label = str(row[1].value or "").strip()
            if not code or not label:
                continue
            values = {}
            for size, col in columns.items():
                cell = sheet.cell(row[0].row, col)
                evaluation = 'formula' if cell.data_type == 'f' else 'direct'
                try:
                    value, issue = cell_value(sheet, cell.coordinate), None
                except (ValueError, SyntaxError, TypeError, ZeroDivisionError):
                    value, issue = None, "Unresolved numeric value or unsupported formula"
                    if cell.data_type == 'f':
                        try:
                            value = number(cached[sheet.title][cell.coordinate].value)
                            issue = 'Cached Excel result; confirm source review before use'
                            evaluation = 'cached'
                        except (ValueError, TypeError):
                            pass
                values[size] = {
                    "value": value,
                    "raw": cell.value,
                    "formula": cell.value if cell.data_type == "f" else None,
                    "cell": cell.coordinate,
                    "issue": issue,
                    "override": False,
                    "evaluation": evaluation,
                }
            key, mapping_status, mapped_from = map_code(code)
            result.append(
                {
                    "key": key,
                    "mapping_status": mapping_status,
                    "mapped_from": mapped_from,
                    "code": code,
                    "label": label.lstrip("* "),
                    "original_label": label,
                    "parser_version": 'xlsx_v2',
                    "unit": "unconfirmed",
                    "confidence": round(sum(1 for cell in values.values() if cell["value"] is not None) / max(len(values), 1), 3),
                    "tolerance": sheet.cell(row[0].row, tolerance_col).value if tolerance_col else None,
                    "source": filename,
                    "sheet": sheet.title,
                    "row": row[0].row,
                    "values": values,
                }
            )
    wb.close()
    cached.close()
    if not result:
        raise ValueError("No measurement rows with size headers found")
    return result

